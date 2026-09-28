import os
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.fonts import addMapping
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
)


class ExportError(RuntimeError):
    """Raised when a meeting report cannot be exported."""


# ---------------------------------------------------------------------------
# Transcript source
# ---------------------------------------------------------------------------

def get_transcript_segments(meeting_data: dict) -> list:
    """
    Current /meetings/analyze format: meeting_data["diarization"]["segments"].
    Older saved results: meeting_data["transcript"]["segments"].
    """
    for key in ("diarization", "transcript"):
        section = meeting_data.get(key) or {}
        if isinstance(section, dict):
            segments = section.get("segments") or []
            if segments:
                return segments
    return []


def get_transcript_full_text(meeting_data: dict) -> str:
    for key in ("diarization", "transcript"):
        section = meeting_data.get(key) or {}
        if isinstance(section, dict) and section.get("full_text"):
            return section["full_text"]
    return ""


def format_segment_line(index: int, segment: dict) -> str:
    return (
        f'SEG_{index:03d} '
        f'[{segment.get("start", "")} - {segment.get("end", "")}] '
        f'{segment.get("speaker") or "Unknown"}: '
        f'{segment.get("text") or ""}'
    )


# ---------------------------------------------------------------------------
# PDF fonts
# ---------------------------------------------------------------------------

# Letters the report must be able to render (Russian + Kazakh-specific).
REQUIRED_GLYPHS = (
    "АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ"
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюя"
    "ӘәҒғҚқҢңӨөҰұҮүҺһІі"
)

PDF_FONT_NAME = "MeetingSans"
PDF_BOLD_FONT_NAME = "MeetingSans-Bold"

# (regular file, bold file) in order of preference.
FONT_CANDIDATES = [
    ("arial.ttf", "arialbd.ttf"),
    ("Arial.ttf", "Arial Bold.ttf"),
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
    ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf"),
    ("NotoSans-Regular.ttf", "NotoSans-Bold.ttf"),
    ("segoeui.ttf", "segoeuib.ttf"),
    ("tahoma.ttf", "tahomabd.ttf"),
    ("FreeSans.ttf", "FreeSansBold.ttf"),
]


def _font_directories() -> list:
    directories = []
    windows_dir = os.environ.get("WINDIR") or os.environ.get("SystemRoot")
    if windows_dir:
        directories.append(os.path.join(windows_dir, "Fonts"))
    directories.append(r"C:\Windows\Fonts")
    local_app_data = os.environ.get("LOCALAPPDATA")
    if local_app_data:
        directories.append(
            os.path.join(local_app_data, "Microsoft", "Windows", "Fonts")
        )
    directories += [
        "/Library/Fonts",
        "/System/Library/Fonts/Supplemental",
        os.path.expanduser("~/Library/Fonts"),
        "/usr/share/fonts",
        "/usr/local/share/fonts",
        os.path.expanduser("~/.fonts"),
        os.path.expanduser("~/.local/share/fonts"),
    ]
    # matplotlib bundles DejaVu Sans; use it if installed.
    try:
        import matplotlib
        directories.append(
            os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
        )
    except Exception:
        pass
    return [d for d in directories if os.path.isdir(d)]


def _find_font_file(filename: str, directories: list):
    for directory in directories:
        direct = os.path.join(directory, filename)
        if os.path.isfile(direct):
            return direct
    # Linux keeps fonts in sub-folders (e.g. /usr/share/fonts/truetype/dejavu).
    for directory in directories:
        for root, _dirs, files in os.walk(directory):
            if filename in files:
                return os.path.join(root, filename)
    return None


def _load_font(name: str, path: str):
    """Load a TTF and confirm it covers every required glyph."""
    try:
        font = TTFont(name, path)
    except Exception:
        return None
    char_map = getattr(font.face, "charToGlyph", {}) or {}
    if any(ord(char) not in char_map for char in REQUIRED_GLYPHS):
        return None
    return font


_registered_fonts = None


def register_pdf_fonts():
    """
    Fallback order: Arial, DejaVu Sans, then other Unicode fonts.
    MEETING_PDF_FONT / MEETING_PDF_BOLD_FONT env vars override the search.
    Returns (regular_font_name, bold_font_name, regular_path).
    """
    global _registered_fonts
    if _registered_fonts:
        return _registered_fonts

    candidates = []
    override = os.environ.get("MEETING_PDF_FONT")
    if override:
        candidates.append((override, os.environ.get("MEETING_PDF_BOLD_FONT")))

    directories = _font_directories()
    for regular_file, bold_file in FONT_CANDIDATES:
        regular_path = _find_font_file(regular_file, directories)
        if regular_path:
            candidates.append(
                (regular_path, _find_font_file(bold_file, directories))
            )

    for regular_path, bold_path in candidates:
        if not regular_path or not os.path.isfile(regular_path):
            continue
        regular = _load_font(PDF_FONT_NAME, regular_path)
        if regular is None:
            continue
        bold = None
        if bold_path and os.path.isfile(bold_path):
            bold = _load_font(PDF_BOLD_FONT_NAME, bold_path)
        if bold is None:
            # Fall back to the regular face so bold text never loses glyphs.
            bold = _load_font(PDF_BOLD_FONT_NAME, regular_path)

        pdfmetrics.registerFont(regular)
        pdfmetrics.registerFont(bold)
        # Make <b> inside Paragraph markup use the Unicode bold face.
        addMapping(PDF_FONT_NAME, 0, 0, PDF_FONT_NAME)
        addMapping(PDF_FONT_NAME, 1, 0, PDF_BOLD_FONT_NAME)
        addMapping(PDF_FONT_NAME, 0, 1, PDF_FONT_NAME)
        addMapping(PDF_FONT_NAME, 1, 1, PDF_BOLD_FONT_NAME)
        _registered_fonts = (PDF_FONT_NAME, PDF_BOLD_FONT_NAME, regular_path)
        return _registered_fonts

    raise ExportError(
        "PDF export failed: no Unicode font with Cyrillic and Kazakh glyphs "
        "was found (tried Arial, DejaVu Sans, Liberation Sans, Noto Sans, "
        "Segoe UI, Tahoma, FreeSans). Install DejaVu Sans or set "
        "MEETING_PDF_FONT to a .ttf file path."
    )


def pdf_text(value) -> str:
    """Escape dynamic text for ReportLab Paragraph markup."""
    if value is None:
        return ""
    return escape(str(value)).replace("\n", "<br/>")


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def export_meeting_docx(meeting_data: dict, output_path: str):
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    document = Document()
    document.styles["Normal"].font.size = Pt(11)

    title = document.add_heading("Meeting Minutes", level=0)
    title.alignment = 1
    document.add_paragraph("Automatically generated meeting report")

    document.add_heading("Meeting Summary", level=1)
    document.add_paragraph(meeting_data.get("summary") or "")

    document.add_heading("Action Items", level=1)
    for index, action in enumerate(meeting_data.get("action_items", []) or [], start=1):
        document.add_heading(f"Action {index}", level=2)
        document.add_paragraph(action.get("task") or "")
        document.add_paragraph(f'Owner: {action.get("owner") or "Not specified"}')
        document.add_paragraph(f'Deadline: {action.get("deadline") or "Not specified"}')
        document.add_paragraph(
            "Evidence segment IDs: "
            + ", ".join(action.get("evidence_segment_ids", []) or [])
        )
        document.add_paragraph(action.get("evidence") or "")

    document.add_heading("Transcript", level=1)
    segments = get_transcript_segments(meeting_data)
    if segments:
        for index, segment in enumerate(segments):
            document.add_paragraph(format_segment_line(index, segment))
    else:
        document.add_paragraph(get_transcript_full_text(meeting_data))

    document.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def export_meeting_pdf(meeting_data: dict, output_path: str):
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    font_name, bold_font_name, _font_path = register_pdf_fonts()

    document = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "MeetingTitle", parent=styles["Title"], fontName=bold_font_name,
        fontSize=20, leading=24, alignment=TA_CENTER, spaceAfter=18,
    )
    heading_style = ParagraphStyle(
        "MeetingHeading", parent=styles["Heading1"], fontName=bold_font_name,
        fontSize=14, leading=18, spaceBefore=12, spaceAfter=8,
    )
    action_heading_style = ParagraphStyle(
        "ActionHeading", parent=styles["Heading2"], fontName=bold_font_name,
        fontSize=11, leading=14, spaceBefore=8, spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "MeetingBody", parent=styles["BodyText"], fontName=font_name,
        fontSize=9, leading=13, spaceAfter=6,
    )

    story = []
    story.append(Paragraph("Meeting Minutes", title_style))
    story.append(Paragraph("Automatically generated meeting report", body_style))

    story.append(Paragraph("Meeting Summary", heading_style))
    story.append(Paragraph(pdf_text(meeting_data.get("summary")), body_style))

    story.append(Paragraph("Action Items", heading_style))
    for index, action in enumerate(meeting_data.get("action_items", []) or [], start=1):
        story.append(Paragraph(f"Action {index}", action_heading_style))
        story.append(Paragraph(
            f"<b>Task:</b> {pdf_text(action.get('task') or 'Not specified')}",
            body_style,
        ))
        story.append(Paragraph(
            f"<b>Owner:</b> {pdf_text(action.get('owner') or 'Not specified')}",
            body_style,
        ))
        story.append(Paragraph(
            f"<b>Deadline:</b> {pdf_text(action.get('deadline') or 'Not specified')}",
            body_style,
        ))
        evidence_ids = action.get("evidence_segment_ids", []) or []
        if evidence_ids:
            story.append(Paragraph(
                f"<b>Evidence:</b> {pdf_text(', '.join(evidence_ids))}",
                body_style,
            ))

    story.append(PageBreak())
    story.append(Paragraph("Transcript", heading_style))

    segments = get_transcript_segments(meeting_data)
    if segments:
        for index, segment in enumerate(segments):
            speaker = segment.get("speaker") or "Unknown"
            text = segment.get("text") or ""
            story.append(Paragraph(
                f"<b>SEG_{index:03d} {pdf_text(speaker)}:</b> {pdf_text(text)}",
                body_style,
            ))
    else:
        full_text = get_transcript_full_text(meeting_data)
        story.append(Paragraph(pdf_text(full_text), body_style))

    story.append(Spacer(1, 6))
    document.build(story)
    return output_path
